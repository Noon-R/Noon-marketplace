#!/usr/bin/env python3
"""Build a .pptx deck from an outline JSON, optionally on top of a template.

Usage:
    python build_pptx.py outline.json -o deck.pptx [-t template.pptx]

The template's slide master / layouts drive all styling. Slides are
addressed to layouts by name (or index), and content is poured into the
layout's placeholders. See references/outline-schema.md for the JSON format.
"""
import argparse
import json
import sys

try:
    from pptx import Presentation
    from pptx.util import Cm, Pt
    from pptx.enum.shapes import PP_PLACEHOLDER
except ImportError:
    sys.exit("python-pptx is required: pip install python-pptx")

R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

BODY_TYPES = (
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.SUBTITLE,
    PP_PLACEHOLDER.VERTICAL_BODY,
    PP_PLACEHOLDER.VERTICAL_OBJECT,
)


def delete_all_slides(prs: Presentation) -> None:
    """Drop sample slides shipped inside the template file."""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.get(R_ID))
        sldIdLst.remove(sldId)


def find_layout(prs: Presentation, ref):
    """Resolve a layout by name (case-insensitive) or by integer index."""
    layouts = [l for m in prs.slide_masters for l in m.slide_layouts]
    if isinstance(ref, int):
        if 0 <= ref < len(layouts):
            return layouts[ref]
        sys.exit(f"layout index {ref} out of range (0-{len(layouts) - 1})")
    for layout in layouts:
        if layout.name == ref:
            return layout
    for layout in layouts:
        if layout.name.lower() == str(ref).lower():
            return layout
    names = ", ".join(f'"{l.name}"' for l in layouts)
    sys.exit(f'layout "{ref}" not found. Available: {names}')


def set_bullets(text_frame, items) -> None:
    """items: list of strings or {"text": ..., "level": n, "bold": bool}."""
    text_frame.clear()
    for i, item in enumerate(items):
        if isinstance(item, str):
            item = {"text": item}
        para = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        para.text = str(item.get("text", ""))
        para.level = int(item.get("level", 0))
        if item.get("bold"):
            for run in para.runs:
                run.font.bold = True
        if item.get("size_pt"):
            for run in para.runs:
                run.font.size = Pt(item["size_pt"])


def get_placeholder(slide, key):
    """Resolve a placeholder by idx (int / numeric str) or by shape name."""
    try:
        return slide.placeholders[int(key)]
    except (ValueError, KeyError):
        pass
    for ph in slide.placeholders:
        if ph.name == key:
            return ph
    available = ", ".join(
        f"idx={p.placeholder_format.idx}({p.name})" for p in slide.placeholders
    )
    sys.exit(f'placeholder "{key}" not found on slide. Available: {available}')


def fill_placeholder(ph, value) -> None:
    if isinstance(value, list):
        set_bullets(ph.text_frame, value)
    else:
        ph.text_frame.text = str(value)


def first_body_placeholder(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.type in BODY_TYPES:
            return ph
    return None


def add_image(slide, spec) -> None:
    path = spec["path"]
    if "placeholder" in spec:
        ph = get_placeholder(slide, spec["placeholder"])
        ph.insert_picture(path)
        return
    kwargs = {}
    if "width_cm" in spec:
        kwargs["width"] = Cm(spec["width_cm"])
    if "height_cm" in spec:
        kwargs["height"] = Cm(spec["height_cm"])
    slide.shapes.add_picture(
        path, Cm(spec.get("left_cm", 2)), Cm(spec.get("top_cm", 4)), **kwargs
    )


def add_table(slide, spec) -> None:
    rows = spec["rows"]
    n_rows, n_cols = len(rows), max(len(r) for r in rows)
    if "placeholder" in spec:
        ph = get_placeholder(slide, spec["placeholder"])
        frame = ph.insert_table(n_rows, n_cols)
    else:
        frame = slide.shapes.add_table(
            n_rows, n_cols,
            Cm(spec.get("left_cm", 2)), Cm(spec.get("top_cm", 4)),
            Cm(spec.get("width_cm", 22)), Cm(spec.get("height_cm", 1.0 * n_rows)),
        )
    table = frame.table
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            table.cell(r, c).text = str(cell)


def build_slide(prs: Presentation, spec: dict) -> None:
    layout = find_layout(prs, spec.get("layout", 1))
    slide = prs.slides.add_slide(layout)

    if "title" in spec and slide.shapes.title is not None:
        slide.shapes.title.text_frame.text = str(spec["title"])

    if "body" in spec:
        ph = first_body_placeholder(slide)
        if ph is None:
            sys.exit(f'layout "{layout.name}" has no body placeholder for "body"')
        fill_placeholder(ph, spec["body"])

    for key, value in spec.get("placeholders", {}).items():
        fill_placeholder(get_placeholder(slide, key), value)

    for img in spec.get("images", []):
        add_image(slide, img)
    for tbl in spec.get("tables", []):
        add_table(slide, tbl)

    # Remove placeholders that ended up empty so no "Click to add text"
    # prompt boxes survive in the output.
    if spec.get("prune_empty", True):
        for ph in list(slide.placeholders):
            if ph.has_text_frame and not ph.text_frame.text.strip():
                if ph.placeholder_format.type != PP_PLACEHOLDER.PICTURE:
                    ph._element.getparent().remove(ph._element)

    if "notes" in spec:
        slide.notes_slide.notes_text_frame.text = str(spec["notes"])


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("outline", help="outline JSON file")
    ap.add_argument("-o", "--output", required=True, help="output .pptx path")
    ap.add_argument("-t", "--template", help="template .pptx (overrides outline)")
    args = ap.parse_args()

    with open(args.outline, encoding="utf-8") as f:
        outline = json.load(f)

    template = args.template or outline.get("template")
    prs = Presentation(template) if template else Presentation()

    if template is None and outline.get("slide_size") == "16:9":
        prs.slide_width = Cm(33.867)
        prs.slide_height = Cm(19.05)

    if not outline.get("keep_existing_slides", False):
        delete_all_slides(prs)

    for spec in outline["slides"]:
        build_slide(prs, spec)

    prs.save(args.output)
    print(f"Saved {len(outline['slides'])} slides -> {args.output}")


if __name__ == "__main__":
    main()
