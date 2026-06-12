#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Inspect a .pptx template: list slide masters, layouts, and placeholders.

Usage:
    python inspect_template.py <template.pptx> [--json]

Run this BEFORE building slides so you know which layout names and
placeholder indices are available in the template.
"""
import argparse
import json
import sys

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("python-pptx is required: pip install python-pptx")


def describe(path: str) -> dict:
    prs = Presentation(path)
    info = {
        "file": path,
        "slide_width_cm": round(Emu(prs.slide_width).cm, 2),
        "slide_height_cm": round(Emu(prs.slide_height).cm, 2),
        "existing_slides": len(prs.slides._sldIdLst),
        "masters": [],
    }
    for m_idx, master in enumerate(prs.slide_masters):
        m = {"index": m_idx, "name": master.name, "layouts": []}
        for l_idx, layout in enumerate(master.slide_layouts):
            lay = {"index": l_idx, "name": layout.name, "placeholders": []}
            for ph in layout.placeholders:
                lay["placeholders"].append({
                    "idx": ph.placeholder_format.idx,
                    "type": str(ph.placeholder_format.type).split(" ")[0],
                    "name": ph.name,
                    "left_cm": round(Emu(ph.left).cm, 2) if ph.left is not None else None,
                    "top_cm": round(Emu(ph.top).cm, 2) if ph.top is not None else None,
                    "width_cm": round(Emu(ph.width).cm, 2) if ph.width is not None else None,
                    "height_cm": round(Emu(ph.height).cm, 2) if ph.height is not None else None,
                })
            m["layouts"].append(lay)
        info["masters"].append(m)
    return info


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("template", help="path to .pptx template")
    ap.add_argument("--json", action="store_true", help="emit raw JSON")
    args = ap.parse_args()

    info = describe(args.template)
    if args.json:
        print(json.dumps(info, ensure_ascii=False, indent=2))
        return

    print(f"File: {info['file']}")
    print(f"Slide size: {info['slide_width_cm']} x {info['slide_height_cm']} cm")
    print(f"Existing slides in file: {info['existing_slides']}")
    for m in info["masters"]:
        print(f"\nMaster [{m['index']}] {m['name']}")
        for lay in m["layouts"]:
            print(f"  Layout [{lay['index']}] \"{lay['name']}\"")
            for ph in lay["placeholders"]:
                print(f"    idx={ph['idx']:<3} type={ph['type']:<12} name=\"{ph['name']}\"")


if __name__ == "__main__":
    main()
